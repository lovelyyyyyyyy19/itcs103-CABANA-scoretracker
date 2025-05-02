import tkinter as tk
from tkinter import messagebox, scrolledtext
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font

FILENAME = "student_scores.xlsx"

# load workbook
def initialize_workbook():
    try:
        wb = openpyxl.load_workbook(FILENAME)
        ws = wb.active
    except FileNotFoundError:
        wb = Workbook()
        ws = wb.active
        ws.append(["Name", "Score", "Remarks"])
        for col in ["A1", "B1", "C1"]:
            ws[col].font = Font(bold=True)
        wb.save(FILENAME)
    return wb, ws

# student data
def add_student_score(name, score):
    wb, ws = initialize_workbook()
    remark = "Pass" if score >= 75 else "Fail"
    ws.append([name, score, remark])
    wb.save(FILENAME)

# Add average to Excel file
def add_average_to_excel():
    wb, ws = initialize_workbook()
    max_row = ws.max_row
    scores = [row[1] for row in ws.iter_rows(min_row=2, max_row=max_row, values_only=True) if isinstance(row[1], (int, float))]

    if scores:
        average = sum(scores) / len(scores)
        
        if ws.cell(row=max_row, column=1).value == "Average":
            ws.delete_rows(max_row, 1)
        ws.append(["Average", average])
        wb.save(FILENAME)

# display records in GUI
def display_all_records():
    wb, ws = initialize_workbook()
    records_text.delete(1.0, tk.END)
    records_text.insert(tk.END, "All Student Records:\n")

    scores = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        name, score, remark = row
        if name is None:
            continue
        records_text.insert(tk.END, f"Name: {name}, Score: {score}, Remarks: {remark}\n")
        if isinstance(score, (int, float)):
            scores.append(score)

    if scores:
        average_score = sum(scores) / len(scores)
        records_text.insert(tk.END, f"\nAverage Score: {average_score:.2f}")
    else:
        records_text.insert(tk.END, "\nAverage Score: N/A")

# On button click
def submit_data():
    name = name_entry.get().strip()
    score_input = score_entry.get().strip()

    if not name or not score_input:
        messagebox.showwarning("Input Error", "Please enter both name and score.")
        return

    try:
        score = float(score_input)
        add_student_score(name, score)
        add_average_to_excel()
        display_all_records()
        name_entry.delete(0, tk.END)
        score_entry.delete(0, tk.END)
        messagebox.showinfo("Success", "Student score added successfully.")
    except ValueError:
        messagebox.showerror("Invalid Input", "Score must be a number.")

# GUI setup
root = tk.Tk()
root.title("Student Score Manager")

tk.Label(root, text="Student Name:").grid(row=0, column=0, padx=10, pady=5, sticky="e")
name_entry = tk.Entry(root, width=30)
name_entry.grid(row=0, column=1, padx=10, pady=5)

tk.Label(root, text="Score:").grid(row=1, column=0, padx=10, pady=5, sticky="e")
score_entry = tk.Entry(root, width=30)
score_entry.grid(row=1, column=1, padx=10, pady=5)

submit_btn = tk.Button(root, text="Add Score", command=submit_data)
submit_btn.grid(row=2, column=0, columnspan=2, pady=10)

records_text = scrolledtext.ScrolledText(root, width=50, height=12)
records_text.grid(row=3, column=0, columnspan=2, padx=10, pady=10)

# Load existing record
display_all_records()

root.mainloop()
